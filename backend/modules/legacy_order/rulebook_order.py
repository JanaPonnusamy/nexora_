"""Rule-book order engine for the legacy Order Process.

Owner decision (2026-10-25): an order can be calculated by the Procurement
rule book instead of the ported VB.NET SQL, while it is still stored in the
SAME OrderNMC tables (OrderManagement / OrderManagementBackup /
OrderHeaderDetails) through order_process.insert_data_into_destination -- so
Qty Check, Bulk Assign, the Procurement screen, exports and WhatsApp all work
unchanged.

The quantity is quantity_profile_t.evaluate_profile_t (the agreed target in
docs/procurement-workspace-unification/06_PROCUREMENT_RULEBOOK.md):

    gross = MAX(avgDaily * maxCover, maxBillQty, maxDayQty, avgBillQty)
    net   = gross - Products.TotalStock - confirmed supplier pending
    packs = CEILING(net / SaleUnit)

with eligibility (valid lines only, no sale in the recency window -> not
ordered, no sale since the last receipt -> suppressed, missing SaleUnit ->
blocked) and the rare-mover 1-pack manual-review line. Parameters are the
Rules Registry's effective values (PRR-QTY-WINDOW / MIN-COVER / MAX-COVER /
RECENCY) for the store, rule-book defaults 90 / 13 / 18 / 15. Confirmed
supplier pending is 0 until PRR-STK-PENDING names its source.

Which engine a store uses is a per-store setting (dbo.LegacyOrderEngine in
OrderNMC, default LEGACY). compare_engines() computes both orders from the
same data without writing anything, for the dry-run review.
"""

from __future__ import annotations

import datetime
import logging
import math

from modules.legacy_order import database, order_process
from modules.procurement.quantity_profile_t import evaluate_profile_t

logger = logging.getLogger(__name__)

ENGINE_LEGACY = "LEGACY"
ENGINE_RULEBOOK = "RULEBOOK"
ENGINES = (ENGINE_LEGACY, ENGINE_RULEBOOK)

RULE_IDS = {
    "window_days": "PRR-QTY-WINDOW",
    "min_days": "PRR-QTY-MIN-COVER",
    "max_days": "PRR-QTY-MAX-COVER",
    "recency_days": "PRR-QTY-RECENCY",
}
DEFAULT_PARAMS = {"window_days": 90, "min_days": 13, "max_days": 18, "recency_days": 15}

# WantedType shown in the Order Workspace / Procurement "Wanted" column: which
# gross term set the quantity.
WANTED = {
    "cover": "Rule Book · Cover Days",
    "max_bill": "Rule Book · Max Bill Qty",
    "max_day": "Rule Book · Max Day Sale",
    "avg_bill": "Rule Book · Avg Bill Qty",
}
WANTED_RARE = "Rule Book · Rare Moving — Manual Review"

_ENGINE_SCHEMA_SQL = """
IF OBJECT_ID('dbo.LegacyOrderEngine', 'U') IS NULL
BEGIN
    CREATE TABLE dbo.LegacyOrderEngine (
        StoreName NVARCHAR(100) NOT NULL PRIMARY KEY,
        Engine VARCHAR(16) NOT NULL DEFAULT 'LEGACY',
        UpdatedAt DATETIME2(0) NOT NULL DEFAULT SYSUTCDATETIME(),
        UpdatedBy NVARCHAR(128) NULL,
        CONSTRAINT CK_LegacyOrderEngine_Engine CHECK (Engine IN ('LEGACY', 'RULEBOOK'))
    );
END;
"""


# ---- per-store engine setting ----------------------------------------------

def _ensure_engine_schema(cur):
    cur.execute(_ENGINE_SCHEMA_SQL)


def get_store_engine(store_name):
    with database.get_central_connection() as conn:
        cur = conn.cursor()
        _ensure_engine_schema(cur)
        conn.commit()
        row = cur.execute(
            "SELECT Engine FROM dbo.LegacyOrderEngine WHERE StoreName = ?", store_name
        ).fetchone()
    return row[0] if row else ENGINE_LEGACY


def set_store_engine(store_name, engine, actor=None):
    engine = str(engine or "").upper()
    if engine not in ENGINES:
        raise ValueError(f"Engine must be one of {', '.join(ENGINES)}.")
    with database.get_central_connection() as conn:
        cur = conn.cursor()
        _ensure_engine_schema(cur)
        cur.execute(
            "MERGE dbo.LegacyOrderEngine AS t USING (SELECT ? AS StoreName) AS s "
            "ON t.StoreName = s.StoreName "
            "WHEN MATCHED THEN UPDATE SET Engine = ?, UpdatedAt = SYSUTCDATETIME(), UpdatedBy = ? "
            "WHEN NOT MATCHED THEN INSERT (StoreName, Engine, UpdatedBy) VALUES (?, ?, ?);",
            store_name, engine, actor, store_name, engine, actor,
        )
        conn.commit()
    return {"store_name": store_name, "engine": engine}


# ---- parameters ------------------------------------------------------------

def rule_params(store_name):
    """The Rules Registry's effective quantity parameters for the store
    (catalog defaults when the store is not on the platform)."""
    params = dict(DEFAULT_PARAMS)
    try:
        from modules.procurement.rules_resolution import REGISTRY
        store_id = database.platform_store_id(store_name)
        tenant_id = database.platform_tenant_id(store_id) if store_id else None
        if store_id and tenant_id:
            values = {r["rule_id"]: r["value"] for r in REGISTRY.resolve(tenant_id, store_id)["rules"]}
            for key, rule_id in RULE_IDS.items():
                if isinstance(values.get(rule_id), int):
                    params[key] = values[rule_id]
    except Exception:
        logger.exception("rule-book params for %s: falling back to defaults", store_name)
    return params


# ---- calculation -----------------------------------------------------------

def _num(value):
    return float(value) if value is not None else 0.0


def evaluate_rows(input_rows, params, as_of):
    """input_rows: dicts from sql/order_rulebook_inputs.sql. Returns one
    decision per product: the input row plus profile T's result."""
    out = []
    for r in input_rows:
        case = {
            "window_sales_qty": _num(r["window_qty"]),
            "bill_frequency": int(r["bill_frequency"] or 0),
            "line_frequency": int(r["line_frequency"] or 0),
            "max_bill_qty": _num(r["max_bill_qty"]),
            "max_day_qty": _num(r["max_day_qty"]),
            "total_stock": _num(r["TotalStock"]),
            "confirmed_pending": 0,
            "sale_unit": None if r["SaleUnit"] is None else float(r["SaleUnit"]),
            "dont_consider": False,
            "last_sale": r["last_sale"],
            "last_received": r["last_received"],
        }
        out.append({**r, "result": evaluate_profile_t(case, as_of, params)})
    return out


def order_rows(decisions, params, store_name, store_code, as_of):
    """The ordered lines as (columns, rows) in the shape the legacy order SQL
    returns, so order_process.insert_data_into_destination stores them."""
    order_id = as_of.strftime("%Y%m%d%H%M%S")
    columns = [src for src, _ in order_process.COLUMN_MAP if src != "OrderQty"] + ["OrderQty"]
    rows = []
    for d in decisions:
        res = d["result"]
        packs = res["packs"]
        if not packs or packs <= 0:
            continue
        window = _num(d["window_qty"])
        avg_daily = window / float(params["window_days"])
        line_freq = int(d["line_frequency"] or 0)
        min_qty = max(math.ceil(avg_daily * params["min_days"]), math.ceil(window / line_freq) if line_freq else 0)
        is_pharma = 0 if d["ProductType"] == 0 else 1
        values = {
            "productcode": d["ProductCode"],
            "productname": d["ProductName"],
            "TotalStock": d["TotalStock"],
            "SaleUnit": d["SaleUnit"],
            "slsqty": window,
            "PurchasePrice": d["PurchasePrice"],
            "MRP": d["MRP"],
            "UnitDescription": d["UnitDescription"],
            "SubLocation": d["SubLocation"],
            "producttype": is_pharma,
            "LastReceivedDate": d["last_received"],
            "LastSaleDate": d["last_sale"],
            "MaxSaleQty": d["max_day_qty"],
            "Transactiondate": d["max_day_date"],
            "WantedDate": as_of,
            "storename": store_name,
            "StoreCode": store_code,
            "OrderId": order_id,
            "MinQty": min_qty,
            "MaxQty": math.ceil(res["gross"] or 0),
            "WantedType": WANTED_RARE if res["reason"] == "RARE_MOVING_MANUAL_REVIEW" else WANTED.get(res["determinant"], "Rule Book"),
            "productTypeName": "Pharma" if is_pharma else "Non Pharma",
            "Frequence": line_freq,
            "Status": 0,
            "OrderQty": packs,
        }
        rows.append(tuple(values[c] for c in columns))
    return columns, rows


def _inputs(store_name, warehouse, params, as_of):
    odata = (as_of - datetime.timedelta(days=params["window_days"])).strftime("%Y-%m-%d")
    sql, values = order_process._bind(order_process._load_sql("order_rulebook_inputs"), {
        "storename": store_name, "ODATA": odata, "warehouse": 1 if warehouse else 0,
    })
    with order_process._connect(database.central_connection_string(), timeout=30) as conn:
        conn.timeout = 0
        cur = conn.cursor().execute(sql, values)
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def build_order(store, warehouse=False, as_of=None):
    """Rule-book order for the store as (columns, rows, summary)."""
    as_of = as_of or datetime.datetime.now()
    params = rule_params(store["store_name"])
    decisions = evaluate_rows(_inputs(store["store_name"], warehouse, params, as_of), params, as_of)
    columns, rows = order_rows(decisions, params, store["store_name"], store["store_code"], as_of)
    reasons = {}
    for d in decisions:
        reasons[d["result"]["reason"]] = reasons.get(d["result"]["reason"], 0) + 1
    return columns, rows, {"params": params, "evaluated": len(decisions), "ordered": len(rows), "reasons": reasons}


# ---- dry run ---------------------------------------------------------------

def compare_engines(store, warehouse, legacy_min_days, legacy_max_days, legacy_recency_days):
    """Both orders from the same central data, nothing written. Returns
    {'params', 'legacy_params', 'rows': per-product comparison, 'summary'}."""
    as_of = datetime.datetime.now()
    params = rule_params(store["store_name"])
    decisions = evaluate_rows(_inputs(store["store_name"], warehouse, params, as_of), params, as_of)
    by_code = {d["ProductCode"]: d for d in decisions}

    odata = (datetime.date.today() - datetime.timedelta(days=90)).strftime("%Y-%m-%d")
    cols, legacy_rows = order_process.fetch_source_data(
        database.central_connection_string(), odata, store["store_name"], store["store_code"],
        legacy_min_days, legacy_max_days, "local", recency_days=legacy_recency_days,
        sql_name="order_warehouse" if warehouse else None,
    )
    idx = {c.lower(): i for i, c in enumerate(cols)}
    legacy = {}
    for r in legacy_rows:
        code = r[idx["productcode"]]
        wanted = r[idx["wantedtype"]]
        # The order process zeroes 'Additional Row' lines after insert.
        qty = 0 if wanted == "Additional Row" else int(r[idx["orderqty"]] or 0)
        legacy[code] = {"name": r[idx["productname"]], "qty": qty, "wanted": wanted,
                        "stock": r[idx["totalstock"]], "maxqty": r[idx["maxqty"]]}

    rows = []
    for code in set(legacy) | {c for c, d in by_code.items() if (d["result"]["packs"] or 0) > 0}:
        d = by_code.get(code)
        lg = legacy.get(code)
        res = d["result"] if d else None
        rb_qty = (res["packs"] or 0) if res else 0
        lg_qty = lg["qty"] if lg else 0
        rows.append({
            "product_code": code,
            "product_name": (d and d["ProductName"]) or (lg and lg["name"]),
            "stock": (d and d["TotalStock"]) if d else (lg and lg["stock"]),
            "sale_unit": d and d["SaleUnit"],
            "window_qty": d and _num(d["window_qty"]),
            "bills": d and d["bill_frequency"],
            "max_bill": d and d["max_bill_qty"],
            "max_day": d and d["max_day_qty"],
            "last_sale": d and d["last_sale"],
            "last_received": d and d["last_received"],
            "legacy_qty": lg_qty,
            "legacy_wanted": lg and lg["wanted"],
            "legacy_max": lg and lg["maxqty"],
            "rulebook_qty": rb_qty,
            "rulebook_reason": res and res["reason"],
            "rulebook_term": res and res["determinant"],
            "rulebook_gross": res and res["gross"] is not None and round(res["gross"], 2),
            "diff": rb_qty - lg_qty,
            "change": ("Only rule book" if rb_qty > 0 and lg_qty <= 0 else
                       "Only legacy" if lg_qty > 0 and rb_qty <= 0 else
                       "Both, same qty" if rb_qty == lg_qty else
                       "Both, more" if rb_qty > lg_qty else "Both, less"),
        })
    rows.sort(key=lambda r: (r["change"], str(r["product_name"] or "")))
    summary = {
        "legacy_lines": sum(1 for r in rows if r["legacy_qty"] > 0),
        "legacy_packs": sum(r["legacy_qty"] for r in rows),
        "rulebook_lines": sum(1 for r in rows if r["rulebook_qty"] > 0),
        "rulebook_packs": sum(r["rulebook_qty"] for r in rows),
    }
    for r in rows:
        summary[r["change"]] = summary.get(r["change"], 0) + 1
    return {
        "store": store["store_name"], "as_of": as_of, "warehouse": warehouse,
        "params": params,
        "legacy_params": {"min_days": legacy_min_days, "max_days": legacy_max_days, "recency_days": legacy_recency_days},
        "rows": rows, "summary": summary,
    }


_COMPARE_COLUMNS = [
    ("change", "Change", 16), ("product_code", "Code", 10), ("product_name", "Product", 34),
    ("stock", "Stock", 8), ("sale_unit", "Pack", 6), ("window_qty", "Sales (window)", 12),
    ("bills", "Bills", 7), ("max_bill", "Max bill", 9), ("max_day", "Max day", 9),
    ("last_sale", "Last sale", 11), ("last_received", "Last GRN", 11),
    ("legacy_qty", "Legacy qty", 10), ("legacy_max", "Legacy max", 10), ("legacy_wanted", "Legacy wanted", 30),
    ("rulebook_qty", "Rule book qty", 12), ("rulebook_gross", "Rule book target", 13),
    ("rulebook_term", "Target set by", 12), ("rulebook_reason", "Rule book result", 30), ("diff", "Diff", 7),
]
_REASON_TEXT = {
    "AUTO": "Ordered",
    "ZERO_NET": "Stock covers the target",
    "REVIEW_RECENCY": "No sale in recency window",
    "SUPPRESSED_NO_SALE_SINCE_RECEIPT": "No sale since last GRN",
    "BLOCKED_SALE_UNIT": "Pack size missing — blocked",
    "RARE_MOVING_MANUAL_REVIEW": "Rare moving — 1 pack, review",
    "EXCLUDED_NO_SALES": "No sales",
}


def compare_workbook(data):
    """compare_engines() result -> .xlsx bytes (Summary + Comparison)."""
    import io
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="1F2937")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    p, lp, s = data["params"], data["legacy_params"], data["summary"]
    lines = [
        (f"{data['store']} — Legacy vs Rule Book (dry run, nothing written)", None),
        ("Calculated", data["as_of"].strftime("%d/%m/%Y %H:%M")),
        ("Demand basis", "Retail D + Transfer-Out TO (warehouse)" if data["warehouse"] else "Retail billing"),
        (None, None),
        ("Rule book parameters", f"{p['window_days']}-day window · {p['recency_days']}-day recency · "
                                 f"{p['min_days']}/{p['max_days']}-day cover · stock = Products.TotalStock"),
        ("Legacy parameters", f"90-day window · {lp['recency_days']}-day recency · {lp['min_days']}/{lp['max_days']}-day cover"),
        (None, None),
        ("Legacy order", f"{s['legacy_lines']} lines · {s['legacy_packs']} packs"),
        ("Rule book order", f"{s['rulebook_lines']} lines · {s['rulebook_packs']} packs"),
        (None, None),
    ]
    for key in ("Only legacy", "Only rule book", "Both, more", "Both, less", "Both, same qty"):
        lines.append((key, s.get(key, 0)))
    for r, (a, b) in enumerate(lines, start=1):
        ws.cell(r, 1, a).font = bold
        if b is not None:
            ws.cell(r, 2, b)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 90

    ws = wb.create_sheet("Comparison")
    for c, (_, label, width) in enumerate(_COMPARE_COLUMNS, start=1):
        cell = ws.cell(1, c, label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        ws.column_dimensions[get_column_letter(c)].width = width
    for r, row in enumerate(data["rows"], start=2):
        for c, (key, _, _) in enumerate(_COMPARE_COLUMNS, start=1):
            value = row.get(key)
            if key == "rulebook_reason":
                value = _REASON_TEXT.get(value, value)
            if isinstance(value, datetime.datetime):
                value = value.date()
            cell = ws.cell(r, c, value)
            if isinstance(value, datetime.date):
                cell.number_format = "DD/MM/YY"
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COMPARE_COLUMNS))}{len(data['rows']) + 1}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

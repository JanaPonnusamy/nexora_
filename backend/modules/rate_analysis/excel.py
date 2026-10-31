"""Generic Excel export for Rate Analysis results.

Writes any ``{title, notes, sheets: [{name, columns, rows}]}`` result: a title
and notes block, a grouped two-row header (columns sharing ``group``), number
formats per column ``format``, per-cell flags from ``row["_flags"]`` (good /
bad / warn) and a bold ``_total`` row.
"""

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

_NUM = {"money": "#,##0.00", "int": "#,##0", "num1": "#,##0.0", "pct": "0.0", "date": "dd-mm-yyyy"}
_FLAG = {"good": PatternFill("solid", fgColor="C6EFCE"), "bad": PatternFill("solid", fgColor="FFC7CE"),
         "warn": PatternFill("solid", fgColor="FFEB9C")}
_HEAD = PatternFill("solid", fgColor="1F4E78")
_GROUP = PatternFill("solid", fgColor="2E75B6")
_TOTAL = PatternFill("solid", fgColor="F2F2F2")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WHITE = Font(bold=True, color="FFFFFF")
_CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def _safe_title(name, used):
    base = "".join(ch for ch in str(name) if ch not in '[]:*?/\\')[:31] or "Sheet"
    title, n = base, 1
    while title in used:
        n += 1
        title = f"{base[:28]}_{n}"
    used.add(title)
    return title


def _write_sheet(ws, sheet, title, notes):
    cols = sheet["columns"]
    ws.cell(1, 1, title).font = Font(bold=True, size=14)
    r = 2
    for note in notes or []:
        ws.cell(r, 1, note).font = Font(italic=True, color="666666")
        r += 1
    r += 1
    grouped = any(c.get("group") for c in cols)
    head = r + 1 if grouped else r
    j = 1
    while j <= len(cols):
        c = cols[j - 1]
        g = c.get("group")
        if grouped and g:
            k = j
            while k < len(cols) and cols[k].get("group") == g:
                k += 1
            cell = ws.cell(r, j, g)
            cell.font, cell.fill, cell.alignment = _WHITE, _GROUP, _CENTER
            if k > j:
                ws.merge_cells(start_row=r, start_column=j, end_row=r, end_column=k)
            for jj in range(j, k + 1):
                ws.cell(r, jj).border = _BORDER
            j = k + 1
        else:
            if grouped:
                ws.merge_cells(start_row=r, start_column=j, end_row=head, end_column=j)
                cell = ws.cell(r, j, c["label"])
                cell.font, cell.fill, cell.alignment, cell.border = _WHITE, _HEAD, _CENTER, _BORDER
            j += 1
    for j, c in enumerate(cols, 1):
        if grouped and not c.get("group"):
            continue
        cell = ws.cell(head, j, c["label"])
        cell.font, cell.fill, cell.alignment, cell.border = _WHITE, _HEAD, _CENTER, _BORDER
    ws.row_dimensions[head].height = 36

    first = head + 1
    rr = first
    for row in sheet["rows"]:
        flags = row.get("_flags") or {}
        total = row.get("_total")
        for j, c in enumerate(cols, 1):
            v = row.get(c["key"])
            if isinstance(v, datetime):
                v = v.date()
            cell = ws.cell(rr, j, v)
            cell.border = _BORDER
            fmt = _NUM.get(c.get("format"))
            if fmt and (isinstance(v, (int, float, date))):
                cell.number_format = fmt
            f = flags.get(c["key"])
            if f in _FLAG:
                cell.fill = _FLAG[f]
            if total:
                cell.font = Font(bold=True)
                if f not in _FLAG:
                    cell.fill = _TOTAL
        rr += 1

    for j, c in enumerate(cols, 1):
        width = c.get("width")
        if not width:
            longest = max([len(str(row.get(c["key"]) or "")) for row in sheet["rows"][:300]] + [6])
            width = min(max(longest + 2, 7), 18)
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.freeze_panes = ws.cell(first, min(4, len(cols)))
    if sheet["rows"]:
        ws.auto_filter.ref = f"A{head}:{get_column_letter(len(cols))}{rr - 1}"


def build_workbook(result):
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sheet in result.get("sheets") or []:
        ws = wb.create_sheet(_safe_title(sheet["name"], used))
        _write_sheet(ws, sheet, result.get("title") or sheet["name"], result.get("notes"))
    if not wb.sheetnames:
        wb.create_sheet("Empty")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
